"""
Purpose: Verify the Step 20 deterministic parser pipeline without live storage or OCR services.
Scope: CSV/Excel normalization, PDF scan detection, split/group metadata, derivative
storage behavior, and Celery task registration.
Dependencies: Parser adapters, worker parse task helpers, pypdf, openpyxl, and storage contracts.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO
from uuid import UUID

from apps.worker.app.celery_app import celery_app
from apps.worker.app.tasks.parse_documents import parse_source_document, store_parse_derivatives
from openpyxl import Workbook  # type: ignore[import-untyped]
from pypdf import PdfWriter
from services.contracts.storage_models import (
    CloseRunStorageScope,
    DerivativeKind,
    ObjectStorageReference,
    StorageBucketKind,
)
from services.jobs.task_names import TaskName
from services.parser.models import ParserSourceDocument
from services.parser.pdf_parser import parse_pdf_document
from services.parser.spreadsheet_parser import parse_spreadsheet_document
from services.storage.checksums import compute_sha256_bytes


def test_csv_parser_normalizes_tables_and_detects_grouped_documents() -> None:
    """Ensure CSV files become normalized tables with split candidates for grouped uploads."""

    source = ParserSourceDocument(
        filename="mixed-support.csv",
        mime_type="text/csv",
        payload=(
            b"Document Type,Reference,Amount\n"
            b"Invoice,INV-001,1200.00\n"
            b"Invoice,INV-002,850.50\n"
            b"Receipt,RCP-001,1200.00\n"
        ),
    )

    result = parse_source_document(source)

    assert result.parser_name == "spreadsheet.csv"
    assert result.tables[0].columns == ("Document Type", "Reference", "Amount")
    assert result.tables[0].rows[0]["Reference"] == "INV-001"
    assert len(result.split_candidates) == 2
    assert result.split_candidates[0].document_type_hint.value == "invoice"
    assert result.normalized_payload() is not None


def test_excel_parser_extracts_each_non_empty_sheet_as_a_table() -> None:
    """Ensure XLSX workbooks produce one normalized table per non-empty worksheet."""

    workbook_payload = _build_workbook_payload()

    result = parse_spreadsheet_document(
        payload=workbook_payload,
        filename="period-support.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert result.parser_name == "spreadsheet.excel"
    assert result.page_count == 2
    assert tuple(table.name for table in result.tables) == ("Invoices", "Receipts")
    assert result.tables[0].rows[0]["Amount"] == "4500"
    assert result.normalized_content_type == "application/json"


def test_pdf_parser_marks_low_text_pdf_as_requiring_ocr_without_running_ocr() -> None:
    """Ensure PDF metadata captures scan-like inputs before OCR routing runs."""

    blank_pdf = _build_blank_pdf_payload()

    result = parse_pdf_document(payload=blank_pdf, filename="scanned-statement.pdf")

    assert result.parser_name == "pdf.pypdf"
    assert result.page_count == 1
    assert result.metadata["requires_ocr"] is True
    assert result.normalized_payload() is not None


def test_parse_derivative_storage_writes_normalized_and_table_payloads() -> None:
    """Ensure parser derivatives are stored with canonical derivative kinds and keys."""

    storage = InMemoryDerivativeStorage()
    result = parse_source_document(
        ParserSourceDocument(
            filename="statement.csv",
            mime_type="text/csv",
            payload=b"Date,Description,Amount\n2026-03-01,Opening balance,1000\n",
        )
    )
    scope = CloseRunStorageScope(
        entity_id=UUID("aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa"),
        close_run_id=UUID("bbbbbbbb-bbbb-bbbb-bbbb-bbbbbbbbbbbb"),
        period_start=date(2026, 3, 1),
        period_end=date(2026, 3, 31),
        close_run_version_no=1,
    )

    derivatives = store_parse_derivatives(
        storage_repository=storage,
        scope=scope,
        document_id=UUID("cccccccc-cccc-cccc-cccc-cccccccccccc"),
        document_version_no=1,
        source_filename="statement.csv",
        parser_result=result,
    )

    assert derivatives.normalized_storage_key is not None
    assert derivatives.extracted_tables_storage_key is not None
    assert storage.derivative_kinds == [
        DerivativeKind.NORMALIZED_DOCUMENT,
        DerivativeKind.EXTRACTED_TABLES,
    ]


def test_document_parse_task_is_registered_on_celery_app() -> None:
    """Ensure worker bootstrap imports the parser task into the canonical Celery app."""

    assert TaskName.DOCUMENT_PARSE_AND_EXTRACT.value in celery_app.tasks


@dataclass(frozen=True)
class InMemoryStorageReference:
    """Describe a stored derivative object for parser integration tests."""

    reference: ObjectStorageReference


class InMemoryDerivativeStorage:
    """Capture derivative writes without requiring MinIO."""

    def __init__(self) -> None:
        """Initialize captured derivative metadata."""

        self.objects: dict[str, bytes] = {}
        self.derivative_kinds: list[DerivativeKind] = []

    def download_source_document(self, *, storage_key: str) -> bytes:
        """Return source bytes when tests exercise the full storage protocol."""

        return self.objects[storage_key]

    def store_derivative(
        self,
        *,
        scope: CloseRunStorageScope,
        document_id: UUID,
        document_version_no: int,
        derivative_kind: DerivativeKind,
        filename: str,
        payload: bytes,
        content_type: str,
        expected_sha256: str | None = None,
    ) -> InMemoryStorageReference:
        """Store derivative bytes under a deterministic in-memory key."""

        del scope, content_type
        checksum = compute_sha256_bytes(payload)
        assert expected_sha256 in {None, checksum}
        object_key = (
            f"documents/derivatives/{document_id}/versions/{document_version_no}/"
            f"{derivative_kind.value}/{filename}"
        )
        self.objects[object_key] = payload
        self.derivative_kinds.append(derivative_kind)
        return InMemoryStorageReference(
            reference=ObjectStorageReference(
                bucket_kind=StorageBucketKind.DERIVATIVES,
                bucket_name="derivatives-bucket",
                object_key=object_key,
            )
        )

    def store_ocr_text(
        self,
        *,
        scope: CloseRunStorageScope,
        document_id: UUID,
        document_version_no: int,
        source_filename: str,
        text: str,
        content_type: str = "text/plain; charset=utf-8",
        expected_sha256: str | None = None,
    ) -> InMemoryStorageReference:
        """Store OCR text under a deterministic in-memory key."""

        del scope, content_type
        payload = text.encode("utf-8")
        checksum = compute_sha256_bytes(payload)
        assert expected_sha256 in {None, checksum}
        object_key = (
            f"documents/ocr/{document_id}/versions/{document_version_no}/{source_filename}.txt"
        )
        self.objects[object_key] = payload
        return InMemoryStorageReference(
            reference=ObjectStorageReference(
                bucket_kind=StorageBucketKind.DERIVATIVES,
                bucket_name="derivatives-bucket",
                object_key=object_key,
            )
        )


def _build_workbook_payload() -> bytes:
    """Create a minimal XLSX workbook fixture in memory."""

    workbook = Workbook()
    invoice_sheet = workbook.active
    invoice_sheet.title = "Invoices"
    invoice_sheet.append(("Vendor", "Amount"))
    invoice_sheet.append(("Vendor A", 4500))
    receipt_sheet = workbook.create_sheet(title="Receipts")
    receipt_sheet.append(("Reference", "Amount"))
    receipt_sheet.append(("RCP-1", 200))

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_blank_pdf_payload() -> bytes:
    """Create a minimal one-page PDF fixture in memory."""

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()
